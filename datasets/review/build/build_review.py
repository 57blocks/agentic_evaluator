"""Builds the answer-key review workbook for claim-denials-snf-v2 and kyc-ubo-v2.

Each item row says what the item asks, shows the answer key in Chinese
columns and the steps that produced it (from explain.mjs, which replays the
generators), and has the cells a domain reviewer fills: a verdict on the
answer key, a right/wrong cell per field, and, for each strong model that was
graded wrong, whether that grading stands. Items a strong model got "wrong"
come first, as the likeliest answer-key errors.

    python3 datasets/review/build/build_review.py [YYYY-MM-DD]
"""

import collections, datetime, json, os, re, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter

from intro import intro_sheet

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "../../tasks"))
DATE = sys.argv[1] if len(sys.argv) > 1 else datetime.date.today().isoformat()
OUT = os.path.normpath(os.path.join(HERE, f"../dataset-review-{DATE}.xlsx"))

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
KEY_HEAD = PatternFill("solid", fgColor="375623")
MODEL_HEAD = PatternFill("solid", fgColor="7F6000")
FILL = PatternFill("solid", fgColor="FFFF00")  # reviewer input
EXAMPLE_FILL = PatternFill("solid", fgColor="DDEBF7")  # pre-filled example rows
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")  # answer key and its derivation
GREY = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
ROW_LINE_PX = 12
MAX_ROW_HEIGHT = 409
ANSWER_LIMIT = 3000
GRADE_LIMIT = 200

VERDICTS = ["正确", "答案有误", "答案不完整", "题目信息不足", "题目有歧义", "不确定"]
FIELD_MARKS = ["对", "错"]
MODEL_MARKS = ["模型确实错", "模型也对（判分太严）", "模型对，标准答案错"]
MODELS8 = ["gpt-5.6-sol", "opus-5.5", "kimi-k3", "sonnet-5", "deepseek-v4.1-flash", "glm-5.3", "qwen3.8-max", "deepseek-v4-pro"]

CLAIM_GLOSS = {
    "root_cause": {"timely_filing": "超过申报期限", "missing_or_invalid_information": "信息缺失或无效", "duplicate_claim": "重复理赔",
                   "benefit_limit_exhausted": "福利天数用尽", "qualifying_stay_not_met": "未满足 3 天住院前提"},
    "correctable": {"yes": "能挽回", "no": "不能挽回", "conditional": "看情况"},
    "next_action": {"corrected_claim": "提交新 / 更正理赔", "appeal": "申诉 / redetermination", "write_off": "核销", "bill_patient": "向患者收费"},
}
CLAIM_FIELDS = [("root_cause", "拒付原因"), ("correctable", "能否挽回"), ("next_action", "下一步"), ("deadline", "期限")]

DATASETS = [
    {
        "sheet": "claim-denials", "task": "claim-denials-snf-v2", "rules_sheet": "claim-denials 规则",
        "runs": ["claim-denials-snf-v2-2026-10-20T02-54-50-815Z", "claim-denials-snf-v2-2026-10-20T07-42-08-589Z"],
        "top": ["gpt-5.6-sol", "opus-5.5", "kimi-k3"],
        "rules": {"duplicate-paid": "R6", "late-filing": "R1", "unprocessable": "R3", "filing-through-date": "R1, R2",
                  "new-benefit-period": "R5, R2", "benefit-exhausted": "R5", "qualifying-stay": "R4, R2"},
        "key_cols": [f"标准答案：{label}" for _, label in CLAIM_FIELDS],
        "field_cols": [f"{label} 对吗" for _, label in CLAIM_FIELDS],
        "free_cols": [],
        "examples": [
            {"id": "cd2-001", "verdict": "正确", "fields": ["对", "对", "对", "对"], "models": {},
             "why": "拒付代码 18 = 重复理赔；记录显示原理赔 2025-03-21 已付款，2025-03-26 又原样提交 → 核销，无期限。推导 3 步都成立。三个模型都答对，模型那几格不用看。"},
            {"id": "cd2-028", "verdict": "正确", "fields": ["对", "对", "对", "对"],
             "models": {"gpt-5.6-sol": "模型确实错", "kimi-k3": "模型确实错"},
             "why": "住院 2025-02-24 到 02-27 共 3 天，观察期不算，前提满足 → 申诉。期限：04-13 + 125 天 = 08-16（4 月剩 17 天、5 月 31、6 月 30、7 月 31、8 月 16）。gpt 和 kimi 答 08-18，多算了 2 天，是模型错，所以它们的「判分成立吗」选「模型确实错」；opus 答对是灰格。"},
        ],
    },
    {
        "sheet": "kyc-ubo-v2", "task": "kyc-ubo-v2", "rules_sheet": "kyc 规则",
        "runs": ["kyc-ubo-v2-2026-10-20T02-54-50-806Z", "kyc-ubo-v2-2026-10-20T07-42-08-592Z"],
        "top": ["gpt-5.6-sol", "opus-5.5"],
        "rules": {"normal": "R1", "nominal-boundary": "R1", "aggregation": "R2, R3", "two-vehicles": "R4, R5",
                  "voting-rights": "R1", "own-shares": "R6", "kg": "R7", "no-control": "R5, R8"},
        "key_cols": ["标准答案：受益所有人", "标准答案：虚拟受益所有人"],
        "field_cols": ["受益所有人名单 对吗", "虚拟受益所有人 对吗"],
        "free_cols": ["漏了谁（应算而没算）", "多了谁（不该算却算了）"],
        "examples": [
            {"id": "ubo2-001", "verdict": "正确", "fields": ["对", "对"], "models": {},
             "why": "Ida Seidel 45%、Mara Hartmann 35%，都超过 25%；Anna Buchholz 20% 不算。Ida Seidel 同时是 Prokurist，不影响她作为股东被算入。两个模型都答对。"},
            {"id": "ubo2-024", "verdict": "正确", "fields": ["对", "对"], "models": {"gpt-5.6-sol": "模型确实错"},
             "why": "公司自有份额 20.000 € 要从基数里扣掉（R6），基数 30.000 €，三人各 10.000 € = 33.3%，都算。gpt 没扣自有份额，按 50.000 € 算成每人 20%，判成无人达标、填了虚拟受益所有人 → 模型确实错。"},
        ],
    },
]
TAIL_COLS = ["正确答案应为（结论不是「正确」时填）", "理由 / 依据", "审阅人", "日期"]


def claim_key(key):
    out = []
    for field, _ in CLAIM_FIELDS:
        v = key[field]
        if field == "deadline":
            out.append(v or "无期限")
        else:
            out.append(f"{v}（{CLAIM_GLOSS[field].get(v, '')}）")
    return out


def kyc_key(key):
    owners = "\n".join(f"{o['name']}：{o['basis']}" for o in key["beneficial_owners"]) or "（无）"
    return [owners, "、".join(key["fictitious_beneficial_owners"]) or "（无）"]


def derivations():
    out = subprocess.run(["node", os.path.join(HERE, "explain.mjs")], check=True, capture_output=True, text=True).stdout
    return json.loads(out)


def rule_titles(task):
    titles = {}
    for line in open(f"{ROOT}/{task}/RULES.md", encoding="utf-8"):
        m = re.match(r"## (R\d+) · (.+)", line)
        if m:
            titles[m.group(1)] = m.group(2).strip()
    return titles


def results(task, runs):
    res = collections.defaultdict(dict)
    for run in runs:
        for line in open(f"{ROOT}/{task}/runs/{run}/evaluations.jsonl"):
            e = json.loads(line)
            if "check" in e["evaluator"] and e["state"] in ("pass", "fail"):
                s = e["subject"]
                res[s["input"]][s["candidate"]] = (e["state"] == "pass", e.get("reason") or "", run)
    return res


def answer_of(task, run, cand, inp):
    p = f"{ROOT}/{task}/runs/{run}/checks/{cand}__{inp}__t0/output.txt"
    if not os.path.exists(p):
        return "(无答案：超时)"
    t = open(p, encoding="utf-8").read().strip()
    m = re.search(r"\{.*\}", t, re.S)
    try:
        t = json.dumps(json.loads(m.group(0)), ensure_ascii=False, indent=1) if m else t
    except ValueError:
        pass
    return t if len(t) <= ANSWER_LIMIT else t[:ANSWER_LIMIT] + "…（截断）"


def item_rows(d, explained):
    task, top = d["task"], d["top"]
    manifest = json.load(open(f"{ROOT}/{task}/checks/manifest.json"))
    res = results(task, d["runs"])
    titles = rule_titles(task)
    top_fails = lambda i: sum(1 for m in top if m in res[i] and not res[i][m][0])
    rows = []
    for i in sorted(manifest, key=lambda i: (top_fails(i) == 0, -top_fails(i), i)):
        key = json.load(open(f"{ROOT}/{task}/references/{i}.txt"))
        ex = explained[i]
        models = {}
        for m in top:
            if m not in res[i]:
                models[m] = {"answer": "(无答案：超时)", "grade": "超时", "passed": None}
                continue
            passed, reason, run = res[i][m]
            models[m] = {"answer": answer_of(task, run, m, i), "grade": "✔ 答对" if passed else "✘ " + reason[:GRADE_LIMIT], "passed": passed}
        rows.append({
            "id": i, "archetype": manifest[i]["archetype"], "split": manifest[i]["split"],
            "question": ex["question"], "input": open(f"{ROOT}/{task}/inputs/{i}.txt", encoding="utf-8").read().strip(),
            "key": claim_key(key) if task.startswith("claim") else kyc_key(key),
            "key_json": json.dumps(key, ensure_ascii=False, indent=1),
            "steps": "\n".join(f"{n}. {s}" for n, s in enumerate(ex["steps"], 1)),
            "rules": "\n".join(f"{x} · {titles.get(x, '')}" for x in d["rules"][manifest[i]["archetype"]].split(", ")),
            "passed": sum(1 for m in MODELS8 if m in res[i] and res[i][m][0]),
            "answered": sum(1 for m in MODELS8 if m in res[i]),
            "top_fails": top_fails(i),
            "models": models,
        })
    return rows


def header_cell(ws, col, text, width, fill=HEAD):
    cell = ws.cell(row=1, column=col, value=text)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = BOX
    ws.column_dimensions[get_column_letter(col)].width = width


def item_sheet(wb, d, rows):
    """Writes one dataset's items; returns where each reviewer column landed, for 汇总."""
    ws = wb.create_sheet(d["sheet"])
    top = d["top"]
    layout = (
        [("优先级", 7, HEAD), ("题号", 10, HEAD), ("题型", 15, HEAD), ("划分", 8, HEAD), ("这题问什么", 30, HEAD),
         ("输入（模型看到的全部内容）", 58, HEAD)]
        + [(c, 22, KEY_HEAD) for c in d["key_cols"]]
        + [("标准答案（原始 JSON，模型要输出的格式）", 30, KEY_HEAD)]
        + [("标准答案是怎么推出来的（逐步核对）", 62, KEY_HEAD),
           ("相关规则（点击跳转）", 26, HEAD), ("8 个模型答对数（分母不含超时）", 10, HEAD), ("最强模型中答错的个数", 9, HEAD)]
        + [("结论：标准答案（下拉选）", 14, HEAD)]
        + [(c, 11, HEAD) for c in d["field_cols"]]
        + [(c, 20, HEAD) for c in d["free_cols"]]
        + [(c, w, HEAD) for c, w in zip(TAIL_COLS, [28, 32, 10, 11])]
    )
    for m in top:
        layout += [(f"{m} 的答案（原始输出）", 40, MODEL_HEAD), (f"{m} 判分", 24, MODEL_HEAD), (f"{m} 判分成立吗（只在答错时选）", 18, MODEL_HEAD)]
    for c, (text, width, fill) in enumerate(layout, 1):
        header_cell(ws, c, text, width, fill)
    col = {text: c for c, (text, _, _) in enumerate(layout, 1)}
    review_first = col["结论：标准答案（下拉选）"]
    review_last = col["日期"]

    by_id = {x["id"]: x for x in rows}
    examples = d["examples"]
    first = 2 + len(examples)
    for r, ex in enumerate(examples, 2):
        write_item(ws, r, by_id[ex["id"]], col, d, example=ex)
    for r, x in enumerate(rows, first):
        write_item(ws, r, x, col, d)

    last = first + len(rows) - 1
    add_list(ws, VERDICTS, [col["结论：标准答案（下拉选）"]], first, last)
    add_list(ws, FIELD_MARKS, [col[c] for c in d["field_cols"]], first, last)
    marks = DataValidation(type="list", formula1='"' + ",".join(MODEL_MARKS) + '"', allow_blank=True)
    ws.add_data_validation(marks)
    for m in top:
        c = col[f"{m} 判分成立吗（只在答错时选）"]
        for r, x in enumerate(rows, first):
            if x["models"][m]["passed"] is False:
                marks.add(f"{get_column_letter(c)}{r}")

    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 48
    ws.auto_filter.ref = f"A1:{get_column_letter(len(layout))}{last}"
    link_rules(ws, [by_id[ex["id"]] for ex in examples] + rows, col["相关规则（点击跳转）"], d)
    return {"verdict": col["结论：标准答案（下拉选）"], "fields": [(c, col[c]) for c in d["field_cols"]],
            "models": [(m, col[f"{m} 判分成立吗（只在答错时选）"]) for m in top], "first": first, "last": last, "n": len(rows)}


def write_item(ws, r, x, col, d, example=None):
    """One item row. An example row comes pre-filled, in blue, and is left out of 汇总 and the dropdowns."""
    top = d["top"]
    priority = "示例" if example else ("高" if x["top_fails"] else "")
    values = [priority, x["id"], x["archetype"], x["split"], x["question"], x["input"],
              *x["key"], x["key_json"], x["steps"], x["rules"], f"{x['passed']} / {x['answered']}", x["top_fails"]]
    for c, v in enumerate(values, 1):
        cell = style(ws.cell(row=r, column=c, value=v))
        if col[d["key_cols"][0]] <= c <= col["标准答案是怎么推出来的（逐步核对）"]:
            cell.fill = KEY_FILL
    review = {}
    if example:
        review = {"结论：标准答案（下拉选）": example["verdict"], "理由 / 依据": example["why"], "审阅人": "（示例）"}
        review.update(zip(d["field_cols"], example["fields"]))
    for name in ["结论：标准答案（下拉选）", *d["field_cols"], *d["free_cols"], *TAIL_COLS]:
        style(ws.cell(row=r, column=col[name], value=review.get(name))).fill = EXAMPLE_FILL if example else FILL
    ws.cell(row=r, column=col["日期"]).number_format = "yyyy-mm-dd"
    for m in top:
        mx = x["models"][m]
        style(ws.cell(row=r, column=col[f"{m} 的答案（原始输出）"], value=mx["answer"]))
        style(ws.cell(row=r, column=col[f"{m} 判分"], value=mx["grade"]))
        mark = style(ws.cell(row=r, column=col[f"{m} 判分成立吗（只在答错时选）"]))
        if mx["passed"] is False:
            mark.value = example["models"].get(m) if example else None
            mark.fill = EXAMPLE_FILL if example else FILL
        else:
            mark.value = "—（答对）" if mx["passed"] else "—（超时）"
            mark.fill = GREY
    if example:
        ws.cell(row=r, column=1).fill = EXAMPLE_FILL
        ws.cell(row=r, column=1).font = Font(name=FONT, size=9, bold=True)
    lines = max(x["input"].count("\n"), x["steps"].count("\n") + len(x["steps"]) // 45) + 2
    ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, max(120, lines * ROW_LINE_PX))


def style(cell):
    cell.font = Font(name=FONT, size=9)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = BOX
    return cell


def add_list(ws, options, cols, first, last):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for c in cols:
        L = get_column_letter(c)
        dv.add(f"{L}{first}:{L}{last}")


def rule_rows(task):
    """Row of each rule on its rules sheet: the header is row 1, then one row per parse_rules entry."""
    return {rid: r for r, (rid, _, _) in enumerate(parse_rules(task), 2) if rid.startswith("R")}


def link_rules(ws, rows, c, d):
    where = rule_rows(d["task"])
    for r, x in enumerate(rows, 2):
        first = x["rules"].split(" · ")[0]
        if first in where:
            cell = ws.cell(row=r, column=c)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{d['rules_sheet']}'!A{where[first]}")
            cell.font = Font(name=FONT, size=9, color="0563C1", underline="single")


def parse_rules(task):
    """RULES.md as (section, title, body) rows: the preamble first, then one row per ## heading."""
    text = open(f"{ROOT}/{task}/RULES.md", encoding="utf-8").read()
    parts = re.split(r"^## ", text, flags=re.M)
    head = parts[0].split("\n", 1)
    rows = [("总则", head[0].lstrip("# ").strip(), readable(head[1]) if len(head) > 1 else "")]
    for part in parts[1:]:
        heading, _, body = part.partition("\n")
        m = re.match(r"(R\d+) · (.+)", heading.strip())
        rid, title = (m.group(1), m.group(2)) if m else ("", heading.strip())
        rows.append((rid, title, readable(body)))
    return rows


def readable(md):
    """Markdown for a cell: wrapped lines joined, quotes in 「」, tables and lists one line per row."""
    out = []
    for para in re.split(r"\n\s*\n", md.strip()):
        lines = [l.rstrip() for l in para.split("\n")]
        if all(l.startswith(">") for l in lines):
            out.append("「" + " ".join(l.lstrip("> ").strip() for l in lines) + "」")
        elif all(l.startswith("|") for l in lines):
            cells = [l.strip("|").split("|") for l in lines if not re.match(r"^\|[\s\-|:]+\|$", l)]
            out.append("\n".join("・" + " ｜ ".join(c.strip() for c in row) for row in cells))
        elif all(re.match(r"^\s*([-*]|\d+\.)\s", l) or l.startswith("  ") for l in lines):
            items, cur = [], ""
            for l in lines:
                if re.match(r"^\s*([-*]|\d+\.)\s", l):
                    if cur:
                        items.append(cur)
                    cur = "・" + re.sub(r"^\s*([-*]|\d+\.)\s", "", l)
                else:
                    cur += " " + l.strip()
            items.append(cur)
            out.append("\n".join(items))
        else:
            out.append(" ".join(l.strip() for l in lines))
    return "\n\n".join(out).replace("**", "").replace("`", "")


def rules_page(wb, title, task):
    ws = wb.create_sheet(title)
    for c, (h, w) in enumerate(zip(["编号", "规则", "原文引用和说明（引号里的是法规 / 手册原文，后面是出处）"], [8, 36, 110]), 1):
        header_cell(ws, c, h, w)
    for r, (rid, t, body) in enumerate(parse_rules(task), 2):
        for c, v in enumerate([rid, t, body], 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10, bold=(c < 3))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
        ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, 14 * (body.count("\n") + 2 + len(body) // 120))
    ws.freeze_panes = "A2"


README = [
    ("数据集答案键审阅", None),
    ("一句话", "请审「标准答案」对不对，不是给模型打分。每道题都写了这题问什么、标准答案、标准答案是怎么一步步推出来的；你逐步核对推导，哪一步不对就指出来。"),
    ("先看哪一页", "「数据集介绍」页：每个数据集考什么、给模型的原始指令（prompt 原文）、模型要回答的每个字段是什么意思、怎么判分、各题型是什么意思。"),
    ("每一行怎么看", "从左往右：①这题问什么 → ②输入（模型看到的全部材料）→ ③标准答案（绿色，已翻成中文）→ ④标准答案的推导（绿色，按步骤编号，每步都引用了输入里的数字和规则编号）→ 标准答案旁边还有原始 JSON（就是模型要输出的格式）→ 黄色格子是你要填的 → 最右边是最强模型的原始输出、判分和「判分成立吗」。"),
    ("要填的格子", "黄色格子：\n・结论：标准答案（必填，下拉选）\n・每个字段 对 / 错（结论不是「正确」时，标出错在哪个字段）；kyc 另有「漏了谁 / 多了谁」\n・正确答案应为、理由 / 依据、审阅人、日期\n・最右边每个答错的模型有一个「判分成立吗」（见下）\n其他格子不要改。"),
    ("结论怎么选", "正确：标准答案对，推导每一步都成立。\n答案有误：至少一个字段或一个人判错了；在字段格标「错」，在「正确答案应为」写出正确的。\n答案不完整：答案缺了该有的内容，例如 kyc 漏了一个应该算的人、该有期限却写了无期限；在「漏了谁」或「正确答案应为」写出来。\n题目信息不足：输入里缺了判断所需的事实，换谁都答不出唯一答案；在「理由」写缺了什么。\n题目有歧义：材料齐全，但两种答案都说得通；在「理由」说明。\n不确定：拿不准，写下疑问。"),
    ("模型「判分成立吗」", "只在该模型被判错的题才需要选（答对的格子是灰色的）。三个选项：\n・模型确实错：标准答案对，模型错了（最常见，例如期限多算了几天）。\n・模型也对（判分太严）：模型的答案和标准答案不同，但同样说得通，判分脚本不该判错；在「理由」写为什么。\n・模型对，标准答案错：这时请同时把「结论」改成「答案有误」。\n答对的题不用看模型答案。"),
    ("示例行", "每个题目页最上面两行（优先级写「示例」、浅蓝色）是已经填好的示范：一道全对的题，一道标准答案对、但有模型答错的题。照着填就行。这两行不计入汇总；同样的题在下面还有一行，正常审阅即可。"),
    ("先审哪些", "「优先级」标「高」的排在最前：最强的模型（gpt-5.6、opus、kimi）答错的题，最可能是答案键或题目本身有问题。时间有限的话，先审这些，再按题型各抽几道其余的。"),
    ("规则在哪", "标准答案按「规则」页里的规则算出：claim-denials 规则（42 CFR、CMS 手册、X12 代码，逐条引用原文）、kyc 规则（§ 3 GwG、BVA 透明登记册 FAQ）。推导里的 R1、R2 是本数据集自己的规则编号，不是法规条款号；题目页「相关规则」可以点击跳到对应规则。模型作答时看不到这些规则。"),
    ("谁来审", "claim-denials：熟悉 Medicare SNF 理赔和申诉的人。kyc-ubo-v2：熟悉德国反洗钱法（GwG）受益所有人认定的人。"),
    ("保密", "表里有全部标准答案，包括封存的 holdout 题。请只在团队内部传阅，不要外发，也不要把题目或答案贴进任何 AI 工具。所有题目、名字、公司都是合成的。"),
    ("汇总", "「汇总」页自动计数：各结论的题数、每个字段被标「错」的次数、审阅人和标准答案的一致率（正确 / 已审），以及每个模型的「判分成立吗」三种选项各多少。"),
    ("示例（仅示意填写格式，不是真实发现）", "某题标准答案期限 2025-10-22。结论：答案有误 ｜ 期限 对吗：错（其他三个字段：对）｜ 正确答案应为：期限 2025-10-27 ｜ 理由：推导第 4 步回单日写错，应为 2025-06-24；+5 +120 = 2025-10-27 ｜ 审阅人：张三 ｜ 日期：2026-10-22。同一题 gpt-5.6 答 2025-10-27 被判错 → 它的「判分成立吗」选「模型对，标准答案错」。"),
]


def readme(wb):
    ws = wb.active
    ws.title = "说明"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for r, (a, b) in enumerate(README, 1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = Font(name=FONT, bold=True, size=14 if r == 1 else 10)
        ca.alignment = Alignment(wrap_text=True, vertical="top")
        if b:
            cb = ws.cell(row=r, column=2, value=b)
            cb.font = Font(name=FONT, size=10)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * (b.count("\n") + 1 + len(b) // 95)
    for r, (a, _) in enumerate(README, 1):
        if a.startswith("示例"):
            ws.cell(row=r, column=2).fill = FILL


def head_row(ws, r, heads):
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def summary(wb, placed):
    ws = wb.create_sheet("汇总", 2)
    bold = Font(name=FONT, bold=True, size=10)
    plain = Font(name=FONT, size=10)
    rng = lambda sheet, c, p: f"'{sheet}'!{get_column_letter(c)}{p['first']}:{get_column_letter(c)}{p['last']}"
    ws.column_dimensions["A"].width = 26
    for c in range(2, 12):
        ws.column_dimensions[get_column_letter(c)].width = 14

    r = 1
    ws.cell(row=r, column=1, value="一、标准答案的结论").font = Font(name=FONT, bold=True, size=12)
    r += 1
    heads = ["数据集", "题数", "已审", *VERDICTS, "一致率（正确 / 已审）"]
    head_row(ws, r, heads)
    first = r + 1
    for sheet, p in placed:
        r += 1
        v = rng(sheet, p["verdict"], p)
        ws.cell(row=r, column=1, value=sheet).font = plain
        ws.cell(row=r, column=2, value=p["n"]).font = plain
        ws.cell(row=r, column=3, value=f"=COUNTA({v})").font = plain
        for c, verdict in enumerate(VERDICTS, 4):
            ws.cell(row=r, column=c, value=f'=COUNTIF({v},"{verdict}")').font = plain
        ws.cell(row=r, column=len(heads), value=f'=IFERROR(D{r}/C{r},"-")').number_format = "0.0%"
    r += 1
    ws.cell(row=r, column=1, value="合计").font = bold
    for c in range(2, len(heads)):
        L = get_column_letter(c)
        ws.cell(row=r, column=c, value=f"=SUM({L}{first}:{L}{r - 1})").font = bold
    ws.cell(row=r, column=len(heads), value=f'=IFERROR(D{r}/C{r},"-")').number_format = "0.0%"
    r += 1
    ws.cell(row=r, column=1, value="目标：一致率达到 95% 以上，数据集才算可以放心使用。").font = Font(name=FONT, size=9, italic=True)

    r += 2
    ws.cell(row=r, column=1, value="二、每个字段被标「错」的次数").font = Font(name=FONT, bold=True, size=12)
    for sheet, p in placed:
        r += 1
        ws.cell(row=r, column=1, value=sheet).font = bold
        for c, (label, col) in enumerate(p["fields"], 2):
            ws.cell(row=r, column=c, value=label.replace(" 对吗", "")).font = bold
            ws.cell(row=r + 1, column=c, value=f'=COUNTIF({rng(sheet, col, p)},"错")').font = plain
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="三、模型被判错的题：判分成立吗").font = Font(name=FONT, bold=True, size=12)
    r += 1
    head_row(ws, r, ["数据集 · 模型", *MODEL_MARKS, "还没选"])
    for sheet, p in placed:
        for m, col in p["models"]:
            r += 1
            v = rng(sheet, col, p)
            ws.cell(row=r, column=1, value=f"{sheet} · {m}").font = plain
            for c, mark in enumerate(MODEL_MARKS, 2):
                ws.cell(row=r, column=c, value=f'=COUNTIF({v},"{mark}")').font = plain
            ws.cell(row=r, column=len(MODEL_MARKS) + 2, value=f"=COUNTBLANK({v})").font = plain
    r += 1
    ws.cell(row=r, column=1, value="「模型也对（判分太严）」多的话要放宽判分脚本；「模型对，标准答案错」应与上面的「答案有误」对得上。").font = Font(name=FONT, size=9, italic=True)


def main():
    explained = derivations()
    wb = Workbook()
    readme(wb)
    intro_sheet(wb, 1)
    placed, sheets = [], []
    for d in DATASETS:
        rows = item_rows(d, explained)
        p = item_sheet(wb, d, rows)
        placed.append((d["sheet"], p))
        sheets.append((d, rows))
        print(d["sheet"], len(rows), "items,", sum(1 for x in rows if x["top_fails"]), "high priority")
    for d, _ in sheets:
        rules_page(wb, d["rules_sheet"], d["task"])
    summary(wb, placed)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

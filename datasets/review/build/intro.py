"""The 数据集介绍 sheet: what each dataset is, what goes in and out, how it is graded, its item types."""
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

TASKS = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "../../tasks"))


def prompt_text(task, file):
    """The instruction the model gets, with the item's input where {{input}} stands."""
    text = open(f"{TASKS}/{task}/prompts/{file}", encoding="utf-8").read().strip()
    return text.replace("{{input}}", "（这里接上题目页「输入」一栏的内容）")

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
SECTION = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MAX_ROW_HEIGHT = 409  # Excel's limit

DATASETS = [
    {
        "name": "claim-denials（Medicare 护理机构理赔拒付分诊）· 300 题",
        "facts": [
            ("业务场景", "美国专业护理机构（SNF）向 Medicare 传统医保（Part A）申报住院费用，被承包商拒付后，收费专员要判断：为什么被拒、能不能挽回、下一步做什么、最晚哪天前做。这个数据集考模型是否懂 Medicare 的这些规则。"),
            ("模型拿到什么", "一份理赔档案：理赔头信息（服务起止日期、入院日期、承包商收到日期）、承包商的拒付通知（835 电子回单的原始代码，例如 CAS*CO*18 表示重复理赔）、机构自己的记录（提交、付款、住院医嘱等时间线）。不给任何规则或手册，模型要自己知道。"),
            ("模型要回答什么", "一个 JSON：\n・root_cause 拒付原因，10 类之一：timely_filing（超过申报期限）、missing_or_invalid_information（信息缺失或无效）、duplicate_claim（重复理赔）、benefit_limit_exhausted（福利天数用尽）、qualifying_stay_not_met（未满足 3 天住院前提）、medical_necessity、coordination_of_benefits、eligibility_terminated、non_covered_service、coding_error\n・correctable 能否挽回：yes / no / conditional（取决于机构还要核实的事实）\n・next_action 下一步：corrected_claim（提交更正或新理赔）、appeal（申请重新审核）、submit_to_other_payer、verify_then_decide、write_off（核销）、bill_patient（向患者收费）\n・deadline 下一步的最后期限（YYYY-MM-DD）；核销和向患者收费时为 null\n・action_note 一两句说明（不参与脚本判分，由评分模型打分）"),
            ("给模型的原始指令（prompt 原文，英文）", prompt_text("claim-denials-snf-v2", "triage.md")),
            ("怎么判分", "脚本逐字段比对 root_cause、correctable、next_action、deadline 四个字段，必须全对才算通过；deadline 要求日期完全一致。"),
            ("题目怎么来的", "全部合成：生成器按固定随机种子造出理赔（gen/generate.mjs），再按「claim-denials 规则」页的规则算出标准答案（gen/rules.mjs）。人名、机构、编号都是虚构的。普通题 110 道，难题（要懂规则才能答对）190 道。"),
        ],
        "types": [
            ("duplicate-paid", "同一笔理赔已经付过款，又重复提交被拒", 30, "duplicate_claim · no · write_off · 无期限", "R6"),
            ("late-filing", "确实超过了一年的申报期限", 30, "timely_filing · no · write_off · 无期限", "R1"),
            ("unprocessable", "信息缺失被退回（MA130，不可处理，没有申诉权）", 50, "missing_or_invalid_information · yes · corrected_claim · 期限 = 一年申报期限", "R3"),
            ("filing-through-date", "承包商按「开始日期」算申报期限而误拒；按规则应按「结束日期」算，实际没超期", 50, "timely_filing · yes · appeal · 期限 = 通知日 + 125 天", "R1、R2"),
            ("new-benefit-period", "承包商以福利天数用尽拒付；其实已开始新的福利期", 50, "benefit_limit_exhausted · yes · appeal · 期限 = 通知日 + 125 天", "R5、R2"),
            ("benefit-exhausted", "同一福利期内超过第 100 天，确实用尽", 40, "benefit_limit_exhausted · no · bill_patient · 无期限", "R5"),
            ("qualifying-stay", "承包商以未满足 3 天住院前提拒付；从住院医嘱算起其实满足（观察期不算住院）", 50, "qualifying_stay_not_met · yes · appeal · 期限 = 通知日 + 125 天", "R4、R2"),
        ],
        "result": "gpt-5.6-sol 92%、kimi-k3 91%、opus-5.5 90%，其余五个模型 20%–49%。差距几乎全来自 R2：申诉期限从「收到通知」起算 120 天，而收到日期推定为通知日后第 5 天，所以是通知日 + 125 天；多数模型答成 + 120 天。",
    },
    {
        "name": "kyc-ubo-v2（德国企业受益所有人认定）· 300 题",
        "facts": [
            ("业务场景", "德国银行做企业客户尽调（KYC）时，要按《反洗钱法》§ 3 GwG 找出公司的「受益所有人」（wirtschaftlich Berechtigte）：直接或间接持股、或控制超过 25% 的自然人。没有人达标时，登记法定代表人为「虚拟受益所有人」。"),
            ("模型拿到什么", "德文的商业登记摘录：公司名、注册资本、总经理（Geschäftsführer）、经理人（Prokura）、股东名单和出资额（欧元），以及章程里关于表决权的说明；持股公司也有自己的摘录。prompt 只说按 § 3 GwG 和登记机关 BVA 的做法，不给具体规则。"),
            ("模型要回答什么", "一个 JSON：\n・beneficial_owners 受益所有人列表，每人 name（与摘录中写法一致）和 basis（依据，例如「45% 的资本」）\n・fictitious_beneficial_owners 虚拟受益所有人姓名列表（只有在无人达标时才填）\n・reasoning 一段说明（不参与脚本判分）"),
            ("给模型的原始指令（prompt 原文，英文）", prompt_text("kyc-ubo-v2", "ubo.md")),
            ("怎么判分", "脚本只比对两份名单（按姓名，不区分大小写）：受益所有人集合、虚拟受益所有人集合都要完全一致。basis 不判分：一个人可能同时满足好几个依据，登记只需要一个。"),
            ("题目怎么来的", "全部合成：生成器按固定种子造出持股结构和德文摘录（gen/generate.mjs），再按「kyc 规则」页的规则算出答案（gen/solve.mjs）。普通题 100 道，陷阱题 200 道。"),
        ],
        "types": [
            ("normal", "股东直接持股，超过 25% 的就是受益所有人", 100, "直接持股 > 25% 的人", "R1"),
            ("nominal-boundary", "出资额正好卡在 25% 边界，例如 25.000 欧元注册资本中 6.250（正好 25%，不算）对 6.251（算）", 30, "只有严格超过 25% 的人", "R1"),
            ("aggregation", "同一人的直接持股和经由其控制的公司的持股要加总", 30, "加总后超过 25% 的人", "R2、R3"),
            ("two-vehicles", "一人通过两个受控公司持股，要加总；旁边还有一个 50/50 的合资公司，50/50 不构成控制", 30, "加总两个受控公司后超过 25% 的人", "R4、R5"),
            ("voting-rights", "表决权和资本比例不一致（例如有无表决权股份），按表决权超过 25% 也算", 30, "资本或表决权任一超过 25% 的人", "R1"),
            ("own-shares", "公司持有自己的股份，计算比例时要从基数中扣除", 20, "扣除自有股份后超过 25% 的人", "R6"),
            ("kg", "GmbH & Co. KG 结构：控制普通合伙人（Komplementär）的人是受益所有人", 30, "控制普通合伙人的人", "R7"),
            ("no-control", "没有人超过 25% 或形成控制，退回登记总经理为虚拟受益所有人（不是 Prokurist）", 30, "受益所有人为空；虚拟受益所有人 = 总经理", "R5、R8"),
        ],
        "result": "gpt-5.6-sol 97%、opus-5.5 97% 明显领先；kimi-k3 86%、sonnet-5 83%、deepseek-flash 81%、glm-5.3 70%、qwen3.8-max 50%（大量超时）、deepseek-v4-pro 46%。错误集中在 two-vehicles、no-control、kg 和 aggregation。",
    },
]


def intro_sheet(wb, index=1):
    ws = wb.create_sheet("数据集介绍", index)
    widths = [26, 54, 10, 44, 12]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    r = 1
    ws.cell(row=r, column=1, value="两个数据集是什么").font = Font(name=FONT, bold=True, size=14)
    r += 2
    for d in DATASETS:
        cell = ws.cell(row=r, column=1, value=d["name"])
        cell.font = Font(name=FONT, bold=True, size=12, color="FFFFFF")
        for c in range(1, 6):
            ws.cell(row=r, column=c).fill = HEAD
        r += 1
        for k, v in d["facts"]:
            ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            cell = ws.cell(row=r, column=2, value=v)
            is_prompt = k.startswith("给模型的原始指令")
            cell.font = Font(name=FONT, size=9 if is_prompt else 10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            # Prompt lines are short and kept as written; prose wraps at about 75 characters.
            lines = v.count("\n") + 1 if is_prompt else v.count("\n") + 1 + len(v) // 75
            ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, max(30, (12 if is_prompt else 15) * lines))
            r += 1
        for c, h in enumerate(["题型", "含义", "题数", "标准答案的样子", "依据的规则"], 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(name=FONT, bold=True, size=10)
            cell.fill = SECTION
            cell.border = BOX
        r += 1
        for t in d["types"]:
            for c, v in enumerate(t, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = BOX
            ws.row_dimensions[r].height = max(30, 15 * (1 + len(t[1]) // 50))
            r += 1
        ws.cell(row=r, column=1, value="8 个模型的结果").font = Font(name=FONT, bold=True, size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        cell = ws.cell(row=r, column=2, value=d["result"])
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (2 + len(d["result"]) // 75)
        r += 3
    return ws
